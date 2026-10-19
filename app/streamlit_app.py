from pathlib import Path
import re
import html

import joblib
import pandas as pd
import streamlit as st
from pypdf import PdfReader
from docx import Document
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

# ================================================================
# SmartHire: resume classification, job matching, skill gaps,
# personalized learning roadmap, and improved Streamlit UI.
# Save this file as app/streamlit_app.py
# ================================================================

ROOT = Path(__file__).resolve().parent.parent
MODEL_PATH = ROOT / "models" / "resume_classifier_eda.pkl"
JOBS_PATH = ROOT / "data" / "processed" / "cleaned_jobs.csv"
if not JOBS_PATH.exists():
    JOBS_PATH = ROOT / "data" / "sample_jobs.csv"

st.set_page_config(
    page_title="SmartHire | Career Guidance",
    page_icon="💼",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ---------- Modern SmartHire UI styling ----------
st.markdown(
    """
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');
    :root { --bg:#070d19; --panel:#101b2b; --border:rgba(148,163,184,.20); --muted:#aab7ca; --accent:#38bdf8; --violet:#a78bfa; }
    .stApp { background:radial-gradient(ellipse at top left,rgba(56,189,248,.10),transparent 36%),radial-gradient(ellipse at top right,rgba(139,92,246,.10),transparent 32%),#070d19; font-family:'Inter',sans-serif; color:#f8fafc; }
    [data-testid="stHeader"] {background:rgba(7,13,25,.88);}
    .block-container {max-width:1500px;padding-top:1.35rem;padding-bottom:2.5rem;}
    h1,h2,h3 {letter-spacing:-.035em;color:#f8fafc;}
    .hero {position:relative;overflow:hidden;padding:2rem 2.2rem;border-radius:23px;margin-bottom:1rem;background:linear-gradient(100deg,rgba(5,15,35,.97) 0%,rgba(30,41,100,.83) 52%,rgba(76,29,149,.60) 100%),url(https://images.unsplash.com/photo-1464822759023-fed622ff2c3b?auto=format&fit=crop&w=1800&q=85) center/cover;border:1px solid rgba(96,165,250,.38);box-shadow:0 18px 55px rgba(0,0,0,.25);}
    .hero:after {content:"";position:absolute;width:270px;height:270px;right:-75px;top:-110px;border-radius:50%;background:rgba(56,189,248,.22);filter:blur(12px);}
    .hero h1 {margin:0;color:#f8fafc;font-size:clamp(2rem,4vw,3.25rem);font-weight:800;line-height:1.12;}
    .hero p {margin:.8rem 0 0;color:#e0eaff;font-size:1rem;line-height:1.7;}
    .eyebrow {color:#7dd3fc;text-transform:uppercase;font-size:.76rem;font-weight:800;letter-spacing:.17em;margin-bottom:.7rem;}
    .hero-gradient {background:linear-gradient(90deg,#38bdf8,#818cf8,#c084fc);-webkit-background-clip:text;background-clip:text;color:transparent;}
    .hero-pills {display:flex;flex-wrap:wrap;gap:.5rem;margin-top:1rem;position:relative;z-index:1;}
    .hero-pill {display:inline-block;padding:.55rem .8rem;border:1px solid rgba(129,140,248,.34);border-radius:12px;background:rgba(15,23,42,.55);color:#f1f5f9;font-size:.82rem;font-weight:600;}
    .section-note {color:var(--muted);font-size:.9rem;}
    div[data-testid="stMetric"] {background:linear-gradient(145deg,#172033,#101827);border:1px solid var(--border);padding:1rem 1.1rem;border-radius:17px;box-shadow:0 8px 28px rgba(0,0,0,.13);min-height:110px;}
    div[data-testid="stMetricLabel"] {color:#cbd5e1;font-weight:600;}
    div[data-testid="stMetricValue"] {color:#7dd3fc;font-weight:800;}
    .skill-chip {display:inline-block;padding:.4rem .75rem;margin:.22rem .25rem .22rem 0;border-radius:999px;background:rgba(16,185,129,.12);color:#a7f3d0;border:1px solid rgba(52,211,153,.35);font-size:.84rem;font-weight:600;}
    .missing-chip {display:inline-block;padding:.4rem .75rem;margin:.22rem .25rem .22rem 0;border-radius:999px;background:rgba(249,115,22,.12);color:#fed7aa;border:1px solid rgba(251,146,60,.35);font-size:.84rem;font-weight:600;}
    div[data-testid="stFileUploader"] {background:linear-gradient(145deg,rgba(30,41,59,.9),rgba(17,24,39,.95));border:1px dashed rgba(129,140,248,.68);border-radius:17px;padding:1rem;}
    div[data-testid="stTabs"] button {font-weight:650;border-radius:10px;}
    div[data-testid="stExpander"] {background:rgba(16,27,43,.72);border:1px solid var(--border);border-radius:14px;overflow:hidden;}
    div[data-testid="stDataFrame"],div[data-testid="stTable"] {border:1px solid var(--border);border-radius:12px;overflow:hidden;}
    .stButton>button,.stLinkButton>a {border-radius:11px;font-weight:700;border:1px solid rgba(129,140,248,.38);}
    .stButton>button[kind="primary"] {background:linear-gradient(100deg,#38bdf8,#8b5cf6);color:#ffffff;border:0;}
    section[data-testid="stSidebar"] {background:linear-gradient(180deg,#0a1222,#111827 75%,#171338);border-right:1px solid rgba(129,140,248,.22);}
    section[data-testid="stSidebar"] h1,section[data-testid="stSidebar"] h2,section[data-testid="stSidebar"] h3 {color:#f8fafc;}
    hr {border-color:var(--border);}
    .panel-card {height:100%;padding:.85rem .95rem;border-radius:0 0 14px 14px;background:linear-gradient(145deg,#142033,#0e1726);border:1px solid rgba(148,163,184,.18);margin-top:-.35rem;}
    /* Keep landing-page images compact and tightly attached to their cards. */
    div[data-testid="stImage"] {margin-bottom:0 !important;}
    div[data-testid="stImage"] img {display:block;width:100%;height:112px;object-fit:cover;border-radius:12px 12px 0 0;}
    .career-start-image img {height:190px;object-fit:cover;border-radius:14px;}
    div[data-testid="stHorizontalBlock"] {gap: .65rem;}
    
    .panel-card h3 {font-size:1.03rem;margin:.1rem 0 .45rem;color:#f8fafc;}
    .panel-card p {font-size:.86rem;line-height:1.55;color:#bdc9da;margin:0;}
    .step-number {display:inline-flex;width:35px;height:35px;align-items:center;justify-content:center;border-radius:50%;background:linear-gradient(135deg,#38bdf8,#8b5cf6);color:#ffffff;font-weight:800;margin-bottom:.5rem;}
    .quote-banner {padding:1.1rem 1.3rem;border-radius:16px;background:linear-gradient(100deg,#78350f,#7c2d12,#422006);border:1px solid rgba(251,191,36,.45);color:#f8fafc;font-weight:600;}
    @media(max-width:768px){.block-container{padding-top:.8rem}.hero{padding:1.4rem 1.2rem;border-radius:17px}}
    </style>
    """,
    unsafe_allow_html=True,
)



# ---------- Model and dataset ----------
@st.cache_resource
def load_model():
    return joblib.load(MODEL_PATH)


@st.cache_data(show_spinner="Loading job listings...")
def load_jobs():
    if not JOBS_PATH.exists():
        raise FileNotFoundError(
            f"Job dataset not found: {JOBS_PATH}. Run src/clean_jobs.py first."
        )

    data = pd.read_csv(
    JOBS_PATH,
    usecols=lambda col: col in [
        "title",
        "description",
        "company_name",
        "location",
        "company",
        "job_title",
        "job_description",
        "job_skills",
        "skill",
        "job_url",
        "link",
        "url"
    ],
    nrows=20000
).fillna("")
    # Real-world datasets often use company_name rather than company.
    aliases = {
        "company_name": "company",
        "company": "company",
        "job_title": "title",
        "description": "description",
        "job_description": "description",
        "job_skills": "skills",
        "skill": "skills",
        "job_url": "url",
        "link": "url",
    }
    rename = {}
    for source, target in aliases.items():
        if source in data.columns and target not in data.columns:
            rename[source] = target
    data = data.rename(columns=rename)

    for col in ["title", "description", "company", "location", "skills", "url"]:
        if col not in data.columns:
            data[col] = ""
        data[col] = data[col].astype(str).str.strip()

    data = data[data["title"].ne("")].copy()
    data["company"] = data["company"].replace("", "Not specified")
    data["location"] = data["location"].replace("", "Not specified")
    data["description"] = data["description"].replace("", "No description supplied.")

    # Keep the original skills column separately; don't overwrite it with
    # title/description text as this was the source of earlier skill-gap issues.
    data["search_text"] = (
        data["title"] + " " + data["description"] + " " + data["skills"]
    ).map(clean_text)

    # Remove duplicate rows where possible.
    dedupe_cols = ["title", "company", "location", "description"]
    data = data.drop_duplicates(subset=dedupe_cols)
    return data.reset_index(drop=True)


# ---------- Resume text extraction ----------
def extract_resume_text(uploaded_file):
    filename = uploaded_file.name.lower()

    if filename.endswith(".pdf"):
        reader = PdfReader(uploaded_file)
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if filename.endswith(".docx"):
        document = Document(uploaded_file)
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)

    if filename.endswith(".txt"):
        return uploaded_file.getvalue().decode("utf-8", errors="ignore")

    return ""


def clean_text(text):
    text = str(text).lower()
    text = re.sub(r"https?://\S+|www\.\S+", " ", text)
    text = re.sub(r"[\r\n\t]+", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


# Canonical names and common aliases. This is transparent keyword matching,
# not semantic proof that a person is proficient in a skill.
SKILL_ALIASES = {
    "python": ["python", "python3"],
    "java": ["java", "core java"],
    "c++": ["c++", "cpp"],
    "c#": ["c#", "c sharp"],
    "sql": ["sql", "structured query language"],
    "pandas": ["pandas"],
    "numpy": ["numpy"],
    "excel": ["excel", "microsoft excel", "ms excel"],
    "machine learning": ["machine learning", "ml algorithms"],
    "deep learning": ["deep learning"],
    "artificial intelligence": ["artificial intelligence", " ai "],
    "data analysis": ["data analysis", "data analytics", "data analyst"],
    "data visualization": ["data visualization", "data visualisation"],
    "statistics": ["statistics", "statistical analysis"],
    "scikit-learn": ["scikit-learn", "scikit learn", "sklearn"],
    "tensorflow": ["tensorflow"],
    "pytorch": ["pytorch", "py torch"],
    "flask": ["flask"],
    "django": ["django"],
    "streamlit": ["streamlit"],
    "html": ["html", "html5"],
    "css": ["css", "css3"],
    "javascript": ["javascript", "java script"],
    "react": ["react", "react.js", "reactjs"],
    "node.js": ["node.js", "nodejs", "node js"],
    "rest api": ["rest api", "restful api", "rest apis"],
    "spring boot": ["spring boot"],
    "git": ["git"],
    "github": ["github"],
    "docker": ["docker"],
    "kubernetes": ["kubernetes", "k8s"],
    "aws": ["aws", "amazon web services"],
    "azure": ["azure", "microsoft azure"],
    "google cloud": ["google cloud", "gcp"],
    "linux": ["linux"],
    "mongodb": ["mongodb", "mongo db"],
    "mysql": ["mysql"],
    "postgresql": ["postgresql", "postgres"],
    "power bi": ["power bi", "powerbi"],
    "tableau": ["tableau"],
    "nlp": ["nlp", "natural language processing"],
    "computer vision": ["computer vision", "opencv"],
    "spark": ["apache spark", "spark"],
    "hadoop": ["hadoop"],
    "communication": ["communication", "verbal communication", "written communication"],
    "project management": ["project management"],
    "problem solving": ["problem solving", "problem-solving"],
    "data structures": ["data structures", "data structure"],
    "algorithms": ["algorithms", "algorithm design"],
    "oop": ["oop", "object oriented programming", "object-oriented programming"],
    "fastapi": ["fastapi", "fast api"],
    "php": ["php"],
    "c": [" c ", " c programming ", " c language "],
    "powerpoint": ["powerpoint", "power point"],
    "collaboration": ["collaboration", "teamwork", "cross-functional teams"],
    "critical thinking": ["critical thinking"],
    "agile": ["agile", "scrum"],
    "data cleaning": ["data cleaning", "data preprocessing", "data wrangling"],
    "etl": ["etl", "extract transform load"],
    "model evaluation": ["model evaluation", "cross validation", "cross-validation"],
    "feature engineering": ["feature engineering"],
    "deployment": ["model deployment", "deployment", "mlops"],
}


def contains_phrase(text, phrase):
    """Match phrases while avoiding accidental partial-word matches."""
    text = f" {clean_text(text)} "
    phrase = clean_text(phrase)
    if not phrase:
        return False
    # Flexible whitespace, but require non-word boundaries around phrase.
    pattern = r"(?<!\w)" + re.escape(phrase).replace(r"\ ", r"\s+") + r"(?!\w)"
    return re.search(pattern, text) is not None


def detect_skills(text):
    found = []
    for canonical, aliases in SKILL_ALIASES.items():
        if any(contains_phrase(text, alias) for alias in aliases):
            found.append(canonical)
    return found


def get_job_required_skills(job_row):
    """
    Prefer the dataset's explicit skills column if it has usable skill terms.
    Otherwise extract known skill keywords from the title and description.
    """
    explicit = str(job_row.get("skills", "")).strip()
    if explicit and explicit.lower() not in {"nan", "none", "not specified"}:
        found = detect_skills(explicit)
        if found:
            return found

        # If skills is a comma/semicolon-separated source column, preserve
        # its terms as a fallback rather than returning an empty list.
        pieces = re.split(r"[,;|/\n]", explicit)
        pieces = [p.strip() for p in pieces if p.strip()]
        if pieces and len(pieces) <= 30:
            return list(dict.fromkeys(pieces))

    return detect_skills(
        f"{job_row.get('title', '')} {job_row.get('description', '')}"
    )


# ---------- Learning roadmap ----------
ROADMAPS = {
    "python": ["Python syntax, functions, collections", "Practice file handling and exceptions", "Build a small automation or data-processing script"],
    "java": ["Review OOP, collections, and exception handling", "Practice JDBC and unit testing", "Build a small REST API with Spring Boot"],
    "sql": ["Practice SELECT, filtering, joins, and aggregation", "Learn subqueries, CTEs, and window functions", "Solve 20 query exercises using a sample database"],
    "pandas": ["Learn Series, DataFrames, filtering, and grouping", "Practice missing-value handling and merges", "Complete an EDA notebook on a public dataset"],
    "numpy": ["Practice arrays, indexing, and broadcasting", "Learn vectorized operations and aggregation", "Use NumPy to support an EDA project"],
    "excel": ["Practice formulas, sorting, and conditional formatting", "Learn PivotTables, lookup functions, and charts", "Build a small KPI dashboard"],
    "machine learning": ["Study train/test splits, leakage, and baselines", "Practice classification and regression with scikit-learn", "Evaluate a model using suitable metrics and cross-validation"],
    "deep learning": ["Review neural-network basics and activation functions", "Train a small model with PyTorch or TensorFlow", "Document experiments and evaluate on held-out data"],
    "statistics": ["Review distributions, mean, variance, and probability", "Practice confidence intervals and hypothesis tests", "Explain statistical findings from a real dataset"],
    "scikit-learn": ["Practice preprocessing and pipelines", "Train models and compare evaluation metrics", "Save a pipeline and test predictions on new inputs"],
    "data analysis": ["Frame a question and identify useful metrics", "Clean and explore a dataset with Pandas", "Present three findings with charts and caveats"],
    "data visualization": ["Choose chart types that fit each question", "Practice labels, units, and readable layouts", "Create a dashboard with a short written summary"],
    "communication": ["Practice explaining a project in a 60-second summary", "Write concise technical notes and findings", "Present a project and ask for feedback"],
    "git": ["Practice status, add, commit, and log", "Learn branches, merges, and conflict resolution", "Publish a project with a clear README"],
    "github": ["Create a clean repository and README", "Use branches and pull requests", "Add setup instructions and screenshots to a project"],
    "html": ["Review semantic HTML and accessible forms", "Build a responsive page structure", "Create a page for your portfolio project"],
    "css": ["Practice flexbox, grid, and responsive design", "Use spacing and typography consistently", "Style a small portfolio page for mobile and desktop"],
    "javascript": ["Review variables, functions, arrays, and objects", "Practice DOM events and form validation", "Build a small interactive page"],
    "rest api": ["Learn HTTP methods, status codes, and JSON", "Call an API and handle errors", "Build or document a small API endpoint"],
    "aws": ["Learn IAM, compute, storage, and networking basics", "Deploy a small practice application", "Document security and cost considerations"],
    "docker": ["Learn images, containers, and Dockerfiles", "Containerize a small application", "Run the app with environment variables and document setup"],
    "power bi": ["Practice importing and cleaning data", "Build relationships and basic DAX measures", "Create a dashboard with a few decision-focused visuals"],
    "tableau": ["Connect a dataset and create calculated fields", "Build interactive charts and filters", "Publish a dashboard with a short insight summary"],
    "nlp": ["Learn tokenization and text normalization", "Build a TF-IDF text-classification baseline", "Evaluate errors and document model limitations"],
    "feature engineering": ["Inspect data types and missingness", "Create useful derived features without leakage", "Compare validation scores before and after changes"],
    "model evaluation": ["Choose metrics appropriate to the task", "Use a validation set or cross-validation", "Inspect a confusion matrix and error examples"],
    "deployment": ["Separate preprocessing and prediction code", "Package dependencies and configuration", "Deploy a small app and test failure cases"],
}


def roadmap_for(skill):
    key = skill.lower()
    if key in ROADMAPS:
        return ROADMAPS[key]
    return [
        f"Review the fundamentals of {skill} using a beginner tutorial.",
        f"Complete a small hands-on exercise using {skill}.",
        f"Add a short project or documented practice result demonstrating {skill}.",
    ]


# ---------- Warm, premium SmartHire header ----------
st.markdown(
    """
    <div class="hero">
        <div class="eyebrow">✦ AI-powered career guidance</div>
        <h1>Turn Your Resume Into<br><span class="hero-gradient">Real Opportunities</span></h1>
        <p>Analyze your resume &nbsp; • &nbsp; Discover matching jobs &nbsp; • &nbsp; Identify skill gaps &nbsp; • &nbsp; Build your learning roadmap</p>
        <div class="hero-pills">
            <span class="hero-pill">✦ AI Resume Analysis</span>
            <span class="hero-pill">▣ Job Recommendations</span>
            <span class="hero-pill">▥ Skill Gap Analysis</span>
            <span class="hero-pill">▤ Learning Roadmap</span>
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)


if not MODEL_PATH.exists():
    st.error(f"Trained model not found at `{MODEL_PATH}`. Run `python src/train_classifier.py` first.")
    st.stop()

if not JOBS_PATH.exists():
    st.error(f"Cleaned job dataset not found at `{JOBS_PATH}`. Run `python src/clean_jobs.py` first.")
    st.stop()

try:
    model = load_model()
    jobs = load_jobs()
except Exception as exc:
    st.error(f"Could not load the model or job dataset: {exc}")
    st.stop()

# Dashboard metrics: values are calculated from the loaded project assets.
metric1, metric2, metric3, metric4 = st.columns(4)
metric1.metric("Job Listings", f"{len(jobs):,}", help="Listings available in your cleaned job dataset")
model_categories = len(getattr(model, "classes_", []))
metric2.metric("Job Categories", f"{model_categories:,}" if model_categories else "Ready", help="Categories available in the trained classifier")
metric3.metric("Skill Keywords", f"{len(SKILL_ALIASES):,}", help="Recognized skill terms used by the app")
metric4.metric("Learning Plan", "4 Weeks", help="Suggested study routine included in the app")

# A small career image strip makes the landing page feel more visual.
image_col, intro_col = st.columns([0.78, 2.22], gap="small")
with image_col:
    st.markdown('<div class="career-start-image">', unsafe_allow_html=True)
    st.image(
        "https://images.unsplash.com/photo-1521737711867-e3b97375f902?auto=format&fit=crop&w=600&h=360&q=80",
        caption="Build your next opportunity",
        width=260,
    )
    st.markdown("</div>", unsafe_allow_html=True)
with intro_col:
    st.markdown("### Your career journey")
    st.markdown(
        "Upload your resume, discover roles that fit your experience, and turn skill gaps into a practical learning plan."
    )
    st.markdown(
        '<div class="panel-card"><h3>✦ Your next step starts here</h3><p>Get a clearer picture of your skills and explore relevant career directions.</p></div>',
        unsafe_allow_html=True,
    )

step_cols = st.columns(5)
steps = [
    ("1", "Upload Resume", "PDF, DOCX or TXT"),
    ("2", "AI Analysis", "Predict category"),
    ("3", "Match Jobs", "Find relevant roles"),
    ("4", "Skill Gap", "Identify missing skills"),
    ("5", "Learning Plan", "Build a roadmap"),
]
for col, (number, title, detail) in zip(step_cols, steps):
    with col:
        st.markdown(f'<div class="panel-card"><span class="step-number">{number}</span><h3>{title}</h3><p>{detail}</p></div>', unsafe_allow_html=True)

st.markdown("### Explore SmartHire")
feature_cols = st.columns(4)
features = [
    ("📄", "Resume Analysis", "AI-assisted category prediction and skill detection."),
    ("💼", "Matching Jobs", "Rank job listings by resume text similarity."),
    ("📊", "Skill Gap Analysis", "Compare resume skills with a target job."),
    ("🧭", "Learning Roadmap", "Get practical steps for skills to improve."),
]
feature_images = [
    "https://images.unsplash.com/photo-1454165804606-c3d57bc86b40?auto=format&fit=crop&w=480&h=220&q=75",
    "https://images.unsplash.com/photo-1521737711867-e3b97375f902?auto=format&fit=crop&w=480&h=220&q=75&crop=faces",
    "https://images.unsplash.com/photo-1551288049-bebda4e38f71?auto=format&fit=crop&w=480&h=220&q=75",
    "https://images.unsplash.com/photo-1516321318423-f06f85e504b3?auto=format&fit=crop&w=480&h=220&q=75",
]
for col, (icon, title, detail), image_url in zip(feature_cols, features, feature_images):
    with col:
        st.image(image_url, use_container_width=True)
        st.markdown(f'<div class="panel-card"><h3>{icon} {title}</h3><p>{detail}</p></div>', unsafe_allow_html=True)

st.markdown('<div class="quote-banner">✦ “A better career is not a dream, it’s a plan. Let’s build your next step with SmartHire.”</div>', unsafe_allow_html=True)
st.markdown("### Upload your resume")
st.caption("Get personalized career insights. Your resume is processed by this app; do not upload sensitive information you do not want to share.")


with st.sidebar:
    st.markdown("# 🎓 Smart<span style='color:#38bdf8'>Hire</span>", unsafe_allow_html=True)
    st.caption("CAREER GUIDANCE")
    st.divider()

    st.markdown("### Main Menu")
    st.markdown("📊  Resume Analysis")
    st.markdown("🎯  Matching Jobs")
    st.markdown("🧩  Skill Gap Analysis")
    st.markdown("🗺️  Learning Roadmap")

    st.divider()
    st.markdown("### Resources")
    st.write(
        "An educational AI-assisted tool to explore job matches "
        "and plan your professional development."
    )
    st.caption(
        "Similarity scores are text similarity, not hiring probabilities."
    )


uploaded_file = st.file_uploader(
    "Upload your resume",
    type=["pdf", "docx", "txt"],
    help="Use a text-based PDF or a DOCX/TXT file. Scanned PDFs may need OCR.",
)

if uploaded_file:
    try:
        resume_text = extract_resume_text(uploaded_file)
    except Exception as exc:
        st.error(f"Could not read the uploaded resume: {exc}")
        st.stop()

    if not resume_text.strip():
        st.error("No readable text was extracted. Try a text-based PDF, DOCX, or TXT file.")
        st.stop()

    cleaned_resume = clean_text(resume_text)
    resume_skills = detect_skills(resume_text)
    st.success(f"Resume text extracted successfully ({len(resume_text.split()):,} words).")

    tab1, tab2, tab3, tab4 = st.tabs([
        "📊 Resume Analysis", "🎯 Matching Jobs", "🧩 Skill Gap", "🗺️ Learning Roadmap"
    ])

    # ---------- 1. Prediction ----------
    with tab1:
        st.header("1. Predicted Resume Category")
        try:
            predicted_category = model.predict([cleaned_resume])[0]
            st.info(f"Predicted category: **{predicted_category}**")

            if hasattr(model, "predict_proba"):
                probabilities = model.predict_proba([cleaned_resume])[0]
                categories = model.classes_
                ranked = sorted(
                    zip(categories, probabilities),
                    key=lambda item: item[1],
                    reverse=True,
                )[:5]
                cat_df = pd.DataFrame(ranked, columns=["Category", "Model score"])
                cat_df["Model score"] = (cat_df["Model score"] * 100).round(2)
                st.caption("Model scores are estimates from the trained classifier; they are not hiring probabilities.")
                st.dataframe(cat_df, hide_index=True, use_container_width=True)
        except Exception as exc:
            st.warning(f"Category prediction could not be completed: {exc}")

        st.subheader("Skills detected in your resume")
        if resume_skills:
            st.markdown(
                "".join(
                    f'<span class="skill-chip">{html.escape(skill.title())}</span>'
                    for skill in resume_skills
                ),
                unsafe_allow_html=True,
            )
        else:
            st.info("No known skills were detected. Consider adding a clearly labeled Skills section to your resume.")

        with st.expander("Preview extracted resume text"):
            st.text_area("Extracted text", resume_text, height=220)

    # ---------- 2. Job matching ----------
    with tab2:
        st.header("2. Matching Jobs")
        st.write("Jobs are ranked by TF-IDF text similarity between the resume and each job listing.")

        search_limit = st.select_slider(
            "Number of listings to search",
            options=[5000, 10000, 20000],
            value=20000,
            help="A smaller limit can be faster. The app searches the first N rows of the cleaned dataset.",
        )
        try:
            jobs_subset = jobs.head(search_limit).copy()
            vectorizer = TfidfVectorizer(
                stop_words="english",
                max_features=30000,
                ngram_range=(1, 2),
                sublinear_tf=True,
            )
            job_matrix = vectorizer.fit_transform(jobs_subset["search_text"])
            resume_vector = vectorizer.transform([cleaned_resume])
            scores = cosine_similarity(resume_vector, job_matrix).ravel() * 100
            jobs_subset["Similarity score (%)"] = scores
            ranked_jobs = jobs_subset.nlargest(10, "Similarity score (%)").copy()
            ranked_jobs["Similarity score (%)"] = ranked_jobs["Similarity score (%)"].round(2)

            if ranked_jobs.empty:
                st.warning("No job listings were available to rank.")
            else:
                st.dataframe(
                    ranked_jobs[["title", "company", "location", "Similarity score (%)"]].rename(
                        columns={
                            "title": "Job title",
                            "company": "Company",
                            "location": "Location",
                        }
                    ),
                    hide_index=True,
                    use_container_width=True,
                )
                st.caption("Scores measure text similarity only; they do not indicate the chance of being hired.")

                st.subheader("Job descriptions and details")
                for rank, (_, job) in enumerate(ranked_jobs.iterrows(), start=1):
                    title = str(job["title"])
                    company = str(job["company"])
                    location = str(job["location"])
                    score = float(job["Similarity score (%)"])
                    with st.expander(f"{rank}. {title} — {company} ({score:.2f}% match)"):
                        left, right = st.columns([3, 1])
                        with left:
                            st.markdown(f"**Location:** {location}")
                            description = str(job["description"])
                            st.write(description[:5000] + ("…" if len(description) > 5000 else ""))
                        with right:
                            st.metric("Text similarity", f"{score:.2f}%")
                            url = str(job.get("url", "")).strip()
                            if url.startswith(("https://", "http://")):
                                st.link_button("Open job posting", url)
                        skills_for_job = get_job_required_skills(job)
                        st.markdown("**Detected skills in this listing**")
                        st.write(", ".join(skills_for_job) if skills_for_job else "No known skill keywords detected.")
        except Exception as exc:
            st.error(f"Could not calculate job matches: {exc}")

    # ---------- 3. Skill gap ----------
    with tab3:
        st.header("3. Skill-Gap Report")
        st.write("Choose a role to compare the skills mentioned in that listing with keywords found in your resume.")

        role_df = jobs[["title", "company", "location"]].copy()
        role_df["role_label"] = (
            role_df["title"].astype(str) + " — " +
            role_df["company"].astype(str) + " (" +
            role_df["location"].astype(str) + ")"
        )
        role_labels = role_df["role_label"].drop_duplicates().tolist()

        if not role_labels:
            st.warning("No job titles are available for skill-gap analysis.")
            missing_skills = []
        else:
            selected_label = st.selectbox("Choose a target job role", role_labels)
            selected_index = role_df.index[role_df["role_label"] == selected_label][0]
            selected_job = jobs.loc[selected_index]
            required_skills = get_job_required_skills(selected_job)

            # Compare canonical skill terms, not the entire vocabulary detected
            # in the resume. This avoids showing unrelated skills as coverage.
            resume_skill_lookup = {skill.lower() for skill in resume_skills}
            matched_skills = [skill for skill in required_skills if skill.lower() in resume_skill_lookup]
            missing_skills = [skill for skill in required_skills if skill.lower() not in resume_skill_lookup]
            coverage = (len(matched_skills) / len(required_skills) * 100) if required_skills else 0.0

            m1, m2, m3 = st.columns(3)
            m1.metric("Required skills detected", len(required_skills))
            m2.metric("Skills matched", len(matched_skills))
            m3.metric("Skill coverage", f"{coverage:.1f}%")

            col1, col2 = st.columns(2)
            with col1:
                st.markdown("**Skills found in your resume**")
                if matched_skills:
                    st.markdown("".join(
                        f'<span class="skill-chip">{html.escape(skill.title())}</span>'
                        for skill in matched_skills
                    ), unsafe_allow_html=True)
                else:
                    st.write("No required skills matched yet.")
            with col2:
                st.markdown("**Skills to consider improving**")
                if missing_skills:
                    st.markdown("".join(
                        f'<span class="missing-chip">{html.escape(skill.title())}</span>'
                        for skill in missing_skills
                    ), unsafe_allow_html=True)
                else:
                    st.write("No detected skill gaps for this vocabulary.")

            with st.expander("Selected job details"):
                st.markdown(f"**{selected_job['title']} — {selected_job['company']}**")
                st.write(f"Location: {selected_job['location']}")
                st.write(str(selected_job["description"])[:5000])

            if not required_skills:
                st.warning(
                    "The selected listing has no skills from the app's current keyword vocabulary. "
                    "The report may be incomplete; consider expanding SKILL_ALIASES for your dataset."
                )
            st.caption(
                "Coverage is the share of detected required-skill keywords also found in your resume. "
                "It does not assess skill depth, experience, or actual proficiency."
            )

    # ---------- 4. Learning roadmap ----------
    with tab4:
        st.header("4. Personalized Learning Roadmap")
        st.write("Use the skill gaps above to create a practical starting plan.")

        if "missing_skills" not in locals():
            missing_skills = []

        if missing_skills:
            selected_learning_skills = st.multiselect(
                "Choose skills you want to learn first",
                options=missing_skills,
                default=missing_skills[: min(3, len(missing_skills))],
            )
            if selected_learning_skills:
                for idx, skill in enumerate(selected_learning_skills, start=1):
                    st.markdown(f"### {idx}. {skill.title()}")
                    steps = roadmap_for(skill)
                    for step_no, step in enumerate(steps, start=1):
                        st.markdown(f"{step_no}. {step}")
                    st.caption("Suggested practice: spend 30–45 minutes per session and save your work in a small project.")
            else:
                st.info("Select at least one skill to see its roadmap.")
        else:
            st.info(
                "Choose a target role in the Skill Gap tab first. If no gaps are found, "
                "the selected job may not contain recognizable skill keywords."
            )

        st.subheader("Suggested 4-week study routine")
        plan = pd.DataFrame(
            [
                ["Week 1", "Learn fundamentals", "Read documentation and take short notes"],
                ["Week 2", "Guided practice", "Complete small exercises for the selected skill"],
                ["Week 3", "Mini-project", "Apply the skill to a small, realistic task"],
                ["Week 4", "Review and publish", "Test your work, write a README, and reflect on gaps"],
            ],
            columns=["Week", "Focus", "Deliverable"],
        )
        st.dataframe(plan, hide_index=True, use_container_width=True)

else:
    st.info("Upload a resume to activate all four SmartHire features.")
    st.markdown("### What you'll get")
    c1, c2 = st.columns(2)
    with c1:
        st.markdown("**📊 Resume analysis**")
        st.write("Predicted category and detected skills.")
        st.markdown("**🎯 Matching jobs**")
        st.write("Ranked listings with company, location, description, and available posting link.")
    with c2:
        st.markdown("**🧩 Skill-gap report**")
        st.write("Skills found in the resume versus skills detected in a selected job.")
        st.markdown("**🗺️ Learning roadmap**")
        st.write("Suggested practice steps for skills you want to improve.")

st.divider()
st.caption(
    "SmartHire is an educational prototype. Review recommendations yourself. "
    "Similarity and skill-coverage values are heuristic indicators, not hiring decisions "
    "or verified measures of proficiency."
)
